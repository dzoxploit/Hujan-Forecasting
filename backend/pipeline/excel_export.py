"""Ekspor hasil ke Excel dengan format 4-baris-header + pewarnaan (sel 18-21)."""

import numpy as np
import pandas as pd

from . import config


def excel_ready(matrix):
    output = matrix.copy()
    output = output.round(3)
    output = output.fillna(config.MISSING_EXPORT_VALUE)
    return output


def prepare_excel_dataframe(matrix, metadata):

    metadata_dict = metadata.set_index("NAMAPOS").to_dict("index")
    output = excel_ready(matrix)

    header_station, header_agency, header_latitude, header_longitude = [], [], [], []

    for station in output.columns:
        header_station.append(station)
        header_agency.append("-")
        info = metadata_dict.get(station, {})
        header_latitude.append(info.get("Latitude", np.nan))
        header_longitude.append(info.get("Longitude", np.nan))

    result = pd.DataFrame(
        index=range(4 + len(output)),
        columns=["Tanggal"] + list(output.columns),
    )

    result.iloc[0, 1:] = header_station
    result.iloc[1, 1:] = header_agency
    result.iloc[2, 1:] = header_latitude
    result.iloc[3, 1:] = header_longitude

    result.iloc[4:, 0] = output.index
    result.iloc[4:, 1:] = output.values

    return result


def style_workbook(workbook):
    from openpyxl.styles import Font, Alignment, Border, Side

    thin = Side(style="thin")
    data_sheets = ["ASLI_-99", "NORMAL_RATIO", "IDW", "RANDOM_FOREST", "HASIL_TERBAIK"]

    for ws in workbook.worksheets:
        if ws.title in data_sheets:
            ws.freeze_panes = "B5"

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for row in range(1, min(4, ws.max_row) + 1):
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).font = Font(bold=True)

        if ws.title in data_sheets:
            for row in range(5, ws.max_row + 1):
                ws.cell(row=row, column=1).number_format = "dd mmm yyyy"

        for col_cells in ws.columns:
            column_letter = col_cells[0].column_letter
            max_length = 0
            for cell in col_cells:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 25)


def highlight_imputed_cells(ws, original, imputed, fallback_mask=None):
    from openpyxl.styles import PatternFill, Font

    green_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")
    green_font = Font(color="006100")
    amber_fill = PatternFill(fill_type="solid", fgColor="FFEB9C")
    amber_font = Font(color="9C6500")

    for col_index, station in enumerate(original.columns, start=2):
        for row_index, date in enumerate(original.index, start=5):
            original_value = original.at[date, station]
            imputed_value = imputed.at[date, station]

            if (
                pd.isna(original_value)
                and pd.notna(imputed_value)
                and imputed_value != config.MISSING_EXPORT_VALUE
            ):
                cell = ws.cell(row=row_index, column=col_index)
                is_fallback = (
                    fallback_mask is not None
                    and bool(fallback_mask.at[date, station])
                )
                if is_fallback:
                    cell.fill = amber_fill
                    cell.font = amber_font
                else:
                    cell.fill = green_fill
                    cell.font = green_font


def highlight_best_result_cells(ws, original, best, source):
    from openpyxl.styles import PatternFill, Font

    green_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")
    green_font = Font(color="006100")
    blue_fill = PatternFill(fill_type="solid", fgColor="DDEBF7")
    blue_font = Font(color="1F4E78")
    amber_fill = PatternFill(fill_type="solid", fgColor="FFEB9C")
    amber_font = Font(color="9C6500")

    for col_index, station in enumerate(original.columns, start=2):
        for row_index, date in enumerate(original.index, start=5):
            original_value = original.at[date, station]
            best_value = best.at[date, station]

            if (
                pd.isna(original_value)
                and pd.notna(best_value)
                and best_value != config.MISSING_EXPORT_VALUE
            ):
                cell = ws.cell(row=row_index, column=col_index)
                method = source.at[date, station]

                if method == "RANDOM_FOREST":
                    cell.fill = green_fill
                    cell.font = green_font
                elif method in ("NORMAL_RATIO", "IDW"):
                    cell.fill = blue_fill
                    cell.font = blue_font
                elif method == "FALLBACK":
                    cell.fill = amber_fill
                    cell.font = amber_font


def export_excel(
    output_path, original, normal_ratio, idw, random_forest, summary, metadata,
    fallback_mask=None, failure_percentage=None, best_result=None, best_source=None,
):
    asli_excel = prepare_excel_dataframe(original, metadata)
    nr_excel = prepare_excel_dataframe(normal_ratio, metadata)
    idw_excel = prepare_excel_dataframe(idw, metadata)
    rf_excel = prepare_excel_dataframe(random_forest, metadata)
    best_excel = (
        prepare_excel_dataframe(best_result, metadata) if best_result is not None else None
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        asli_excel.to_excel(writer, sheet_name="ASLI_-99", index=False, header=False)
        nr_excel.to_excel(writer, sheet_name="NORMAL_RATIO", index=False, header=False)
        idw_excel.to_excel(writer, sheet_name="IDW", index=False, header=False)
        rf_excel.to_excel(writer, sheet_name="RANDOM_FOREST", index=False, header=False)

        if best_excel is not None:
            best_excel.to_excel(writer, sheet_name="HASIL_TERBAIK", index=False, header=False)

        summary.to_excel(writer, sheet_name="RINGKASAN_POS", index=False)
        metadata.to_excel(writer, sheet_name="METADATA_POS", index=False)

        if failure_percentage is not None and not failure_percentage.empty:
            failure_percentage.to_excel(writer, sheet_name="PERSENTASE_GAGAL", index=False)

        style_workbook(writer.book)

        highlight_imputed_cells(writer.sheets["NORMAL_RATIO"], original, normal_ratio, fallback_mask)
        highlight_imputed_cells(writer.sheets["IDW"], original, idw, fallback_mask)
        highlight_imputed_cells(writer.sheets["RANDOM_FOREST"], original, random_forest, fallback_mask)

        if best_result is not None and best_source is not None and "HASIL_TERBAIK" in writer.sheets:
            highlight_best_result_cells(writer.sheets["HASIL_TERBAIK"], original, best_result, best_source)
