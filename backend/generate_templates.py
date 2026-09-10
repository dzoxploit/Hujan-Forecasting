"""Membuat file Excel template untuk data hujan (dipakai untuk upload di
dashboard atau sebagai contoh bagi pengguna baru).

Format harus persis mengikuti pembaca di `pipeline/data_io.py`:
    Baris 1 = Nama PCH
    Baris 2 = Longitude
    Baris 3 = Latitude
    Baris 4 = Header kolom (bebas, tidak dipakai pembaca)
    Baris 5 dst = Tanggal (kolom A) + curah hujan per PCH (kolom B dst)

Menghasilkan dua file di ../templates/:
    TEMPLATE_KOSONG_HUJAN.xlsx      - kerangka kosong siap diisi
    TEMPLATE_CONTOH_TERISI.xlsx     - contoh terisi (dengan beberapa sel
                                       kosong) untuk uji coba unggah/imputasi
"""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT_DIR = Path(__file__).resolve().parent.parent / "templates"
OUT_DIR.mkdir(parents=True, exist_ok=True)

STATIONS = [
    {"nama": "PCH_CONTOH_1", "lon": 110.1234, "lat": -7.1234},
    {"nama": "PCH_CONTOH_2", "lon": 110.4567, "lat": -7.3456},
    {"nama": "PCH_CONTOH_3", "lon": 110.7891, "lat": -7.5678},
    {"nama": "PCH_CONTOH_4", "lon": 110.2222, "lat": -7.7777},
    {"nama": "PCH_CONTOH_5", "lon": 110.8888, "lat": -7.2222},
]

HEADER_FILL = PatternFill(fill_type="solid", fgColor="DDEBF7")
LABEL_FILL = PatternFill(fill_type="solid", fgColor="F2F2F2")
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")


def _write_header_block(ws, stations):
    ws.cell(row=1, column=1, value="Nama PCH").font = BOLD
    ws.cell(row=2, column=1, value="Longitude").font = BOLD
    ws.cell(row=3, column=1, value="Latitude").font = BOLD
    ws.cell(row=4, column=1, value="Tanggal").font = BOLD

    for col_index, station in enumerate(stations, start=2):
        ws.cell(row=1, column=col_index, value=station["nama"])
        ws.cell(row=2, column=col_index, value=station["lon"])
        ws.cell(row=3, column=col_index, value=station["lat"])
        ws.cell(row=4, column=col_index, value="Curah Hujan (mm)")

    for row in range(1, 5):
        for col in range(1, len(stations) + 2):
            cell = ws.cell(row=row, column=col)
            cell.alignment = CENTER
            cell.fill = HEADER_FILL if row in (1, 4) else LABEL_FILL

    ws.column_dimensions["A"].width = 14
    for col_index in range(2, len(stations) + 2):
        ws.column_dimensions[get_column_letter(col_index)].width = 16

    ws.freeze_panes = "B5"


def build_empty_template(path, stations, n_example_rows=14):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Hujan"
    _write_header_block(ws, stations)

    dates = pd.date_range("2016-01-01", periods=n_example_rows, freq="D")
    for row_offset, date in enumerate(dates):
        row = 5 + row_offset
        date_cell = ws.cell(row=row, column=1, value=date.to_pydatetime())
        date_cell.number_format = "dd/mm/yyyy"
        date_cell.alignment = CENTER
        # Kolom nilai sengaja dibiarkan kosong - siap diisi pengguna.

    note_row = 5 + n_example_rows + 1
    ws.cell(row=note_row, column=1, value=(
        "Catatan: tambahkan baris tanggal berikutnya di bawah ini (satu baris = "
        "satu hari), isi nilai curah hujan (mm) tiap PCH. Kosongkan sel untuk "
        "tanggal yang datanya tidak tersedia - JANGAN diisi 0 (0 berarti hari "
        "tanpa hujan, beda dengan data yang tidak diketahui)."
    )).font = Font(italic=True, size=9, color="898781")

    wb.save(path)


def build_filled_example(path, stations, n_days=120, missing_fraction=0.12, seed=42):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Hujan"
    _write_header_block(ws, stations)

    rng = np.random.default_rng(seed)
    dates = pd.date_range("2016-01-01", periods=n_days, freq="D")

    for row_offset, date in enumerate(dates):
        row = 5 + row_offset
        date_cell = ws.cell(row=row, column=1, value=date.to_pydatetime())
        date_cell.number_format = "dd/mm/yyyy"
        date_cell.alignment = CENTER

        for col_index, station in enumerate(stations, start=2):
            if rng.random() < missing_fraction:
                continue  # sel dibiarkan kosong = data hilang

            # Curah hujan sintetis: sering 0 (tidak hujan), kadang hujan deras.
            if rng.random() < 0.6:
                value = 0.0
            else:
                value = round(float(rng.gamma(shape=2.0, scale=8.0)), 1)

            ws.cell(row=row, column=col_index, value=value).alignment = CENTER

    wb.save(path)


def main():
    empty_path = OUT_DIR / "TEMPLATE_KOSONG_HUJAN.xlsx"
    filled_path = OUT_DIR / "TEMPLATE_CONTOH_TERISI.xlsx"

    build_empty_template(empty_path, STATIONS)
    build_filled_example(filled_path, STATIONS)

    print(f"Tersimpan: {empty_path}")
    print(f"Tersimpan: {filled_path}")


if __name__ == "__main__":
    main()
